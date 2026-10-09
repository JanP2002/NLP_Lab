"""Krok 4: arkusze do anotacji w Google Sheets (xlsx) i import wypełnionych arkuszy.

Generacja (z katalogu Lab1):  python -m src.sheets_io --iteration 1 --guidelines v1
Pliki z annotation/iterN/to_annotate/ wgrywamy na Google Drive, każdy anotator dostaje tylko swój.
Wypełnione arkusze pobieramy jako .xlsx do annotation/iterN/exports/ i czytamy read_annotations().

Układ arkusza „Anotacja”: każde zdanie zajmuje dwa wiersze:
  wiersz „tekst”     – sent_id, etykieta zdania (lista), tokeny w kolumnach t1..tK, uwagi
  wiersz „etykieta”  – pod każdym tokenem lista z etykietą słowa (puste = O)
"""
import argparse
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import config

SHEET = "Anotacja"
FIXED_COLS = ["nr", "sent_id", "etykieta_zdania", "wiersz"]
TEXT_ROW, LABEL_ROW = "tekst", "etykieta"

SENT_COLORS = {"POZ": "C6EFCE", "NEG": "FFC7CE", "NEU": "EDEDED", "AMB": "FFF2CC"}
TOKEN_COLORS = {"WYR_POZ": "C6EFCE", "WYR_NEG": "FFC7CE", "WZM": "FFE0B2", "OSL": "DDEBF7", "ODW": "E4D7F5"}
LEGEND = [
    ("POZ", "zdanie: pozytywna ocena"),
    ("NEG", "zdanie: negatywna ocena"),
    ("NEU", "zdanie: brak oceny (fakt, opis)"),
    ("AMB", "zdanie: równoważne plusy i minusy"),
    ("WYR_POZ", "słowo: niesie pozytywną ocenę (polecam, miły, ok)"),
    ("WYR_NEG", "słowo: niesie negatywną ocenę (okropne, brud, niemiły)"),
    ("WZM", "słowo: wzmacnia ocenę (bardzo, absolutnie)"),
    ("OSL", "słowo: osłabia ocenę (trochę, raczej)"),
    ("ODW", "słowo: odwraca ocenę (nie, brak, bez)"),
]


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


def _list_validation(options: list[str], what: str) -> DataValidation:
    return DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                          showErrorMessage=True, errorTitle="Niepoprawna etykieta",
                          error=f"Wybierz {what} z listy: {', '.join(options)}")


def annotator_order(n: int, iteration: int, annotator: str) -> np.ndarray:
    """Każdy anotator dostaje zdania w innej (powtarzalnej) kolejności."""
    rng = np.random.default_rng([config.SEEDS[iteration], config.ANNOTATORS.index(annotator)])
    return rng.permutation(n)


def build_workbook(sample: pd.DataFrame, annotator: str, iteration: int, guidelines: str) -> Workbook:
    tokens = [t.split() for t in sample["text"]]
    k = max(map(len, tokens))
    tok_headers = [f"t{j}" for j in range(1, k + 1)]
    headers = FIXED_COLS + tok_headers + ["uwagi"]
    first_tok = len(FIXED_COLS) + 1
    last_tok_letter = get_column_letter(first_tok + k - 1)
    uwagi_col = len(headers)

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = _fill("D9D9D9")

    dv_sent = _list_validation(config.SENT_LABELS, "etykietę zdania")
    dv_tok = _list_validation(config.TOKEN_LABELS, "etykietę słowa")
    ws.add_data_validation(dv_sent)
    ws.add_data_validation(dv_tok)
    editable, sep = _fill("FFFDF0"), Border(bottom=Side(style="thin", color="999999"))

    for nr, i in enumerate(annotator_order(len(sample), iteration, annotator), start=1):
        sent_id, toks = int(sample["sent_id"].iat[i]), tokens[i]
        r_text = 2 * nr
        r_lab = r_text + 1
        ws.cell(r_text, 1, nr)
        ws.cell(r_text, 2, sent_id)
        ws.cell(r_text, 3).fill = editable
        ws.cell(r_text, 4, TEXT_ROW)
        ws.cell(r_lab, 4, LABEL_ROW).font = Font(italic=True, color="666666")
        dv_sent.add(ws.cell(r_text, 3).coordinate)
        for j, tok in enumerate(toks):
            cell = ws.cell(r_text, first_tok + j, tok)
            cell.data_type = "s"  # tokeny typu „=)” nie mogą stać się formułą
            ws.cell(r_lab, first_tok + j).fill = editable
        last = get_column_letter(first_tok + len(toks) - 1)
        dv_tok.add(f"{get_column_letter(first_tok)}{r_lab}:{last}{r_lab}")
        ws.cell(r_text, uwagi_col).fill = editable
        for col in range(1, uwagi_col + 1):
            ws.cell(r_lab, col).border = sep

    last_row = 2 * len(sample) + 1
    for label, color in SENT_COLORS.items():
        ws.conditional_formatting.add(f"C2:C{last_row}", CellIsRule(operator="equal", formula=[f'"{label}"'], fill=_fill(color)))
    tok_area = f"{get_column_letter(first_tok)}2:{last_tok_letter}{last_row}"
    for label, color in TOKEN_COLORS.items():
        ws.conditional_formatting.add(tok_area, CellIsRule(operator="equal", formula=[f'"{label}"'], fill=_fill(color)))

    widths = {1: 5, 2: 8, 3: 15, 4: 9, uwagi_col: 40}
    for j in range(k):
        longest = max((len(t[j]) for t in tokens if j < len(t)), default=0)
        widths[first_tok + j] = min(max(longest + 2, 10), 18)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = ws.cell(2, first_tok)

    _add_instructions(wb, annotator, iteration, guidelines, len(sample), last_row)
    return wb


def _add_instructions(wb: Workbook, annotator: str, iteration: int, guidelines: str, n: int, last_row: int):
    ws = wb.create_sheet("Instrukcja", 0)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 95
    rows = [
        ("Zadanie 1: anotacja PolEmo 2.0", None),
        ("Anotator", annotator),
        ("Iteracja", iteration),
        ("Wytyczne", f"wytyczne_{guidelines}.md (przeczytaj przed rozpoczęciem)"),
        ("Postęp", f'=COUNTA({SHEET}!C2:C{last_row})&" / {n} zdań"'),
        (None, None),
        ("Krok 1", f"W arkuszu „{SHEET}” w kolumnie etykieta_zdania wybierz POZ / NEG / NEU / AMB."),
        ("Krok 2", "W wierszu „etykieta” pod słowami wybierz etykiety słów. Puste pole = brak etykiety."),
        ("Krok 3", "Wątpliwości zapisz w kolumnie uwagi."),
        ("Nie zmieniaj", "słów, kolejności wierszy ani kolumn; nie dodawaj wierszy."),
        ("Pracujesz sam", "nie konsultuj zdań, nie sprawdzaj etykiet PolEmo, nie używaj LLM."),
        ("Na koniec", "Plik → Pobierz → Microsoft Excel (.xlsx) i wyślij plik koordynatorowi."),
        (None, None),
        ("Legenda", None),
    ]
    for a, b in rows:
        ws.append([a, b])
    ws["A1"].font = Font(bold=True, size=14)
    ws["A14"].font = Font(bold=True)
    for label, desc in LEGEND:
        ws.append([label, desc])
        ws.cell(ws.max_row, 1).fill = _fill({**SENT_COLORS, **TOKEN_COLORS}[label])
    for row in ws.iter_rows(min_row=2, max_row=12, max_col=1):
        row[0].font = Font(bold=True)
    for row in ws.iter_rows(max_col=2):
        row[1].alignment = Alignment(wrap_text=True)


def generate(iteration: int, guidelines: str) -> list[Path]:
    d = config.iter_dir(iteration)
    sample = pd.read_csv(d / "sample.csv")
    out_dir = d / "to_annotate"
    out_dir.mkdir(parents=True, exist_ok=True)
    (d / "exports").mkdir(exist_ok=True)
    paths = []
    for a in config.ANNOTATORS:
        path = out_dir / f"iter{iteration}_{a}.xlsx"
        build_workbook(sample, a, iteration, guidelines).save(path)
        paths.append(path)
    return paths


# --- Import wypełnionych arkuszy (używany w kroku 6) --------------------------

def _norm_label(v) -> str:
    return str(v).strip().upper() if v is not None and str(v).strip() else config.NO_LABEL


def _norm_token(v) -> str:
    if isinstance(v, float) and v.is_integer():  # Sheets mógł zamienić „2020” na liczbę
        v = int(v)
    return "" if v is None else str(v)


def read_annotations(path: Path, sample: pd.DataFrame, annotator: str) -> tuple[pd.DataFrame, list[str]]:
    """Zwraca (DataFrame: sent_id, annotator, sent_label, token_labels, uwagi) oraz listę problemów."""
    ws = load_workbook(path, data_only=True)[SHEET]
    header = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(header) if h}
    tok_cols = [i for i, h in enumerate(header) if isinstance(h, str) and re.fullmatch(r"t\d+", h)]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    texts = dict(zip(sample["sent_id"], sample["text"]))
    allowed_tok = set(config.TOKEN_LABELS) | {config.NO_LABEL}

    records, issues = [], []
    i = 0
    while i < len(rows):
        r = rows[i]
        if r[col["wiersz"]] != TEXT_ROW:
            i += 1
            continue
        where = f"wiersz {i + 2}"
        if i + 1 >= len(rows) or rows[i + 1][col["wiersz"]] != LABEL_ROW:
            issues.append(f"{where}: brak wiersza „{LABEL_ROW}” pod zdaniem")
            i += 1
            continue
        lab = rows[i + 1]
        sid = int(r[col["sent_id"]])
        if sid not in texts:
            issues.append(f"{where}: sent_id {sid} spoza próbki")
            i += 2
            continue
        toks = texts[sid].split()
        n = len(toks)
        if [_norm_token(r[c]) for c in tok_cols[:n]] != toks:
            issues.append(f"{where}: tokeny zdania {sid} różnią się od próbki (edycja arkusza?)")
        if any(_norm_label(lab[c]) != config.NO_LABEL for c in tok_cols[n:]):
            issues.append(f"{where}: etykiety poza zakresem tokenów zdania {sid}")
        tok_labels = [_norm_label(lab[c]) for c in tok_cols[:n]]
        bad = sorted(set(tok_labels) - allowed_tok)
        if bad:
            issues.append(f"{where}: nieznane etykiety słów {bad} w zdaniu {sid}")
        sent_label = _norm_label(r[col["etykieta_zdania"]])
        if sent_label == config.NO_LABEL:
            issues.append(f"{where}: brak etykiety zdania {sid}")
        elif sent_label not in config.SENT_LABELS:
            issues.append(f"{where}: nieznana etykieta zdania {sent_label!r} ({sid})")
        records.append({"sent_id": sid, "annotator": annotator, "sent_label": sent_label,
                        "token_labels": tok_labels, "uwagi": r[col["uwagi"]] or ""})
        i += 2

    df = pd.DataFrame(records)
    seen = df["sent_id"] if len(df) else pd.Series(dtype=int)
    if seen.duplicated().any():
        issues.append(f"powtórzone sent_id: {sorted(seen[seen.duplicated()].unique())}")
    missing = set(texts) - set(seen)
    if missing:
        issues.append(f"brak {len(missing)} zdań z próbki: {sorted(missing)[:10]}")
    return df, issues


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--iteration", type=int, required=True, choices=sorted(config.SEEDS))
    p.add_argument("--guidelines", required=True, help="wersja wytycznych, np. v1")
    args = p.parse_args()
    for path in generate(args.iteration, args.guidelines):
        print(f"Zapisano {path}")


if __name__ == "__main__":
    main()
