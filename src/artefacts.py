"""Pliki iterN_artefacts/ w tym samym układzie co w iteracji 1.

    python -m src.artefacts --iteration 2

Z eksportu (annotation/iterN/export/) i – jeśli istnieje – z rozstrzygnięcia (annotation/iterN/adjudicated.csv)
tworzy w annotation/iterN_artefacts/:
  disagreements_sentences.csv   zdania, w których A1 != A2 (kolumny jak w iteracji 1)
  comparison_tokens.csv         wszystkie tokeny: A1, A2, zgodność (kolumny jak w iteracji 1)
  analiza_rozbieznosci.xlsx     arkusz analizy rozbieżności (kolumny, listy rozwijane i format jak w iteracji 1).
      Wypełniane automatycznie: sent_id, tekst, etykiety A1 i A2, typ rozbieżności (para etykiet) oraz – gdy jest
      adjudicated.csv – „Decyzja sędziego”. Reszta (przyczyna, uzasadnienie, propozycja zmiany wytycznych)
      to ręczna analiza zespołu. Uwagi superanotatora z to_adjudicate.* trafiają do „Uzasadnienie decyzji”.
Istniejący analiza_rozbieznosci.xlsx nie jest nadpisywany bez --force (może zawierać ręczną pracę).
Wypełnione arkusze iterN_A1.xlsx i iterN_A2.xlsx pochodzą od anotatorów i tu nie są ruszane.
"""
from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd

import config
from src import adjudicate

ANALIZA_COLUMNS = ["sent_id", "Tekst zdania", "Etykieta A1", "Etykieta A2", "Typ rozbieżności",
                   "Przyczyna rozbieżności", "Decyzja sędziego", "Uzasadnienie decyzji",
                   "Proponowana zmiana wytycznych"]
TYPES = ["POZ/NEU", "NEG/NEU", "POZ/NEG", "AMB/POZ", "AMB/NEG", "AMB/NEU", "Tokeny", "Inne"]
PAIR_TYPE = {frozenset(p.split("/")): p for p in TYPES[:6]}


def _write_analiza(df: pd.DataFrame, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Analiza rozbieżności"
    ws.append(ANALIZA_COLUMNS)
    for row in df[ANALIZA_COLUMNS].itertuples(index=False):
        ws.append(["" if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    for col, w in zip("ABCDEFGHI", (14, 45, 16, 16, 25, 38, 24, 40, 42)):
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="244062")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    dv_dec = DataValidation(type="list", formula1='"' + ",".join(config.SENT_LABELS) + '"', allow_blank=True)
    dv_typ = DataValidation(type="list", formula1='"' + ",".join(TYPES) + '"', allow_blank=True,
                            showErrorMessage=False)  # jak w iteracji 1: można wpisać własny opis
    ws.add_data_validation(dv_dec)
    ws.add_data_validation(dv_typ)
    dv_dec.add("G2:G500")
    dv_typ.add("E2:E500")
    wb.save(path)


def build(iteration: int, out_dir: Path | None = None, force: bool = False) -> dict[str, Path]:
    out_dir = out_dir or config.filled_sheets_dir(iteration)
    out_dir.mkdir(parents=True, exist_ok=True)
    exp = config.export_dir(iteration)
    for need in ("sentences_wide.csv", "tokens_wide.csv"):
        if not (exp / need).exists():
            raise FileNotFoundError(f"Brak {exp / need}. Najpierw: python -m src.export --iteration {iteration}")
    sw = pd.read_csv(exp / "sentences_wide.csv", encoding="utf-8-sig")
    tw = pd.read_csv(exp / "tokens_wide.csv", encoding="utf-8-sig")
    paths = {}

    conf = sw[sw["A1"] != sw["A2"]].sort_values("sent_id")
    dis = pd.DataFrame({"sent_id": conf["sent_id"], "sentence_label_A1": conf["A1"],
                        "sentence_label_A2": conf["A2"], "_merge": "both", "zgodnosc": False})
    paths["disagreements"] = out_dir / "disagreements_sentences.csv"
    dis.to_csv(paths["disagreements"], index=False)

    comp = pd.DataFrame({"sent_id": tw["sent_id"], "pozycja_tokenu": tw["tok_idx"], "token": tw["token"],
                         "A1": tw["A1"], "A2": tw["A2"], "zgodnosc": tw["zgodne"]})
    comp = comp.sort_values(["sent_id", "pozycja_tokenu"]).reset_index(drop=True)
    paths["comparison"] = out_dir / "comparison_tokens.csv"
    comp.to_csv(paths["comparison"], index=False)

    # decyzje i uwagi superanotatora (jeśli już są)
    decision, remark = {}, {}
    adj = adjudicate.adjudicated_path(iteration)
    if adj.exists():
        a = pd.read_csv(adj, encoding="utf-8-sig")
        decision = dict(zip(a["sent_id"], a["final"]))
    try:
        sh = adjudicate.read_sheet(iteration, include_analiza=False)
        if "uwagi" in sh.columns:
            remark = {int(i): u for i, u in zip(sh["sent_id"], sh["uwagi"]) if isinstance(u, str) and u.strip()}
    except FileNotFoundError:
        pass

    rows = pd.DataFrame({
        "sent_id": conf["sent_id"].to_numpy(),
        "Tekst zdania": conf["text"].to_numpy(),
        "Etykieta A1": conf["A1"].to_numpy(),
        "Etykieta A2": conf["A2"].to_numpy(),
        "Typ rozbieżności": [PAIR_TYPE.get(frozenset((x, y)), "Inne") for x, y in zip(conf["A1"], conf["A2"])],
        "Przyczyna rozbieżności": "",
        "Decyzja sędziego": [decision.get(i, "") for i in conf["sent_id"]],
        "Uzasadnienie decyzji": [remark.get(int(i), "") for i in conf["sent_id"]],
        "Proponowana zmiana wytycznych": "",
    })
    paths["analiza"] = out_dir / "analiza_rozbieznosci.xlsx"
    if paths["analiza"].exists() and not force:
        print(f"POMINIĘTO {paths['analiza']} – już istnieje (użyj --force, żeby nadpisać; może zawierać ręczną analizę)")
    else:
        _write_analiza(rows, paths["analiza"])

    for k, p in paths.items():
        print(f"{k:14s} -> {p}")
    print(f"iteracja {iteration}: {len(sw)} zdań, konfliktów {len(conf)}, tokenów {len(comp)}, "
          f"decyzji sędziego wpisanych: {sum(1 for v in rows['Decyzja sędziego'] if v)}")
    return paths


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--iteration", type=int, required=True)
    ap.add_argument("--out-dir", type=Path, default=None, help="katalog wyjściowy (domyślnie annotation/iterN_artefacts)")
    ap.add_argument("--force", action="store_true", help="nadpisz istniejący analiza_rozbieznosci.xlsx")
    a = ap.parse_args()
    build(a.iteration, a.out_dir, a.force)


if __name__ == "__main__":
    main()
