"""Testy kroków 1–4: parsowanie, losowanie, arkusze (zapis → wypełnienie → odczyt)."""
import pandas as pd
import pytest
from openpyxl import load_workbook

import config
from src import sampling, sheets_io
from src.data import build_corpus, parse_line


# --- krok 1 -----------------------------------------------------------------

def test_parse_line_uses_double_underscore_separator():
    assert parse_line("Wszystko stare i okropne . __label__z_minus_m") == ("Wszystko stare i okropne .", "z_minus_m")
    with pytest.raises(ValueError):
        parse_line("Zdanie bez etykiety _label_z_zero")


def test_build_corpus_dedupes_and_flags_conflicts():
    lines = ["A b c d e . __label__z_zero", "A b c d e . __label__z_zero",
             "X y z w v . __label__z_plus_m", "X y z w v . __label__z_minus_m", ""]
    df, counts = build_corpus(lines)
    assert len(df) == 2 and counts["duplicate_rows_removed"] == 2
    assert df.set_index("text")["gold_conflict"].to_dict() == {"A b c d e .": False, "X y z w v .": True}
    assert df["sent_id"].tolist() == [0, 2]


def test_build_corpus_rejects_unknown_label():
    with pytest.raises(ValueError):
        build_corpus(["Zdanie . __label__z_plus_s"])


# --- krok 2 -----------------------------------------------------------------

def _toy_corpus(per_class=50):
    rows = []
    for k, label in enumerate(config.SENT_LABELS):
        for j in range(per_class):
            sid = k * 1000 + j
            rows.append((sid, f"zdanie numer {sid} ma kilka słów .", label, False, 7))
    return pd.DataFrame(rows, columns=["sent_id", "text", "gold", "gold_conflict", "n_tokens"])


def test_make_sample_is_stratified_deterministic_and_respects_exclusion():
    df = _toy_corpus()
    excl = set(range(0, 40))  # większość klasy POZ wykluczona
    s1 = sampling.make_sample(df, 40, seed=1, exclude_ids=excl)
    assert s1["gold"].value_counts().to_dict() == {l: 10 for l in config.SENT_LABELS}
    assert not set(s1["sent_id"]) & excl
    assert s1.equals(sampling.make_sample(df, 40, seed=1, exclude_ids=excl))
    assert not s1.equals(sampling.make_sample(df, 40, seed=2, exclude_ids=excl))


def test_real_iteration1_sample_is_balanced_and_excludes_guideline_examples():
    sample = pd.read_csv(config.iter_dir(1) / "sample.csv")
    key = pd.read_csv(config.iter_dir(1) / "_key_gold.csv")
    assert len(sample) == config.N_SAMPLE and sample["sent_id"].is_unique
    assert "gold" not in sample.columns  # anotatorzy nie widzą gold
    assert key["gold"].value_counts().to_dict() == {l: config.N_SAMPLE // 4 for l in config.SENT_LABELS}
    assert not set(sample["sent_id"]) & sampling.excluded_ids(1)
    assert sample["n_tokens"].between(config.MIN_TOKENS, config.MAX_TOKENS).all()


# --- krok 4 -----------------------------------------------------------------

TOY = pd.DataFrame({"sent_id": [11, 22, 33],
                    "text": ["Lekarz bardzo miły , ale nie polecam .",
                             "Wesoło =) w 2020 roku .",
                             "Pokój miał 20 m2 ."]})


def _fill_and_save(tmp_path, edit):
    wb = sheets_io.build_workbook(TOY, "A1", 1, "v1")
    edit(wb[sheets_io.SHEET])
    path = tmp_path / "filled.xlsx"
    wb.save(path)
    return path


def _row_of(ws, sent_id):
    return next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value == sent_id)


def test_roundtrip_reads_labels_back(tmp_path):
    first_tok = len(sheets_io.FIXED_COLS) + 1

    def edit(ws):
        r = _row_of(ws, 11)
        ws.cell(r, 3, "NEG")
        for j, lab in {1: "WZM", 2: "WYR_POZ", 5: "ODW", 6: "WYR_POZ"}.items():
            ws.cell(r + 1, first_tok + j, lab)
        ws.cell(r, ws.max_column, "wątpliwe")
        ws.cell(_row_of(ws, 22), 3, "POZ")
        ws.cell(_row_of(ws, 33), 3, "NEU")

    df, issues = sheets_io.read_annotations(_fill_and_save(tmp_path, edit), TOY, "A1")
    assert issues == []
    row = df.set_index("sent_id").loc[11]
    assert row["sent_label"] == "NEG" and row["uwagi"] == "wątpliwe"
    assert row["token_labels"] == ["O", "WZM", "WYR_POZ", "O", "O", "ODW", "WYR_POZ", "O"]
    assert df.set_index("sent_id").loc[22, "token_labels"] == ["O"] * 6


def test_tokens_like_formulas_and_numbers_stay_text(tmp_path):
    path = _fill_and_save(tmp_path, lambda ws: None)
    ws = load_workbook(path)[sheets_io.SHEET]
    r = _row_of(ws, 22)
    assert ws.cell(r, len(sheets_io.FIXED_COLS) + 2).value == "=)"


def test_read_reports_problems(tmp_path):
    first_tok = len(sheets_io.FIXED_COLS) + 1

    def edit(ws):
        r = _row_of(ws, 33)
        ws.cell(r, first_tok, "Pokoj")              # edycja tokenu
        ws.cell(r + 1, first_tok + 6, "WZM")        # etykieta za końcem zdania (5 tokenów, K=8)
        ws.cell(_row_of(ws, 11), 3, "POZYTYWNE")    # spoza listy

    _, issues = sheets_io.read_annotations(_fill_and_save(tmp_path, edit), TOY, "A1")
    text = "\n".join(issues)
    assert "tokeny zdania 33" in text
    assert "poza zakresem tokenów zdania 33" in text
    assert "nieznana etykieta zdania 'POZYTYWNE'" in text
    assert "brak etykiety zdania 22" in text


def test_annotators_get_different_orders():
    orders = [tuple(sheets_io.annotator_order(100, 1, a)) for a in config.ANNOTATORS]
    assert len(set(orders)) == len(orders)


@pytest.mark.parametrize("annotator", config.ANNOTATORS)
def test_generated_iteration1_sheets_align_with_sample(annotator):
    sample = pd.read_csv(config.iter_dir(1) / "sample.csv")
    path = config.iter_dir(1) / "to_annotate" / f"iter1_{annotator}.xlsx"
    df, issues = sheets_io.read_annotations(path, sample, annotator)
    assert set(df["sent_id"]) == set(sample["sent_id"])
    # pusty arkusz: jedyny problem to brak etykiet zdań
    assert len(issues) == config.N_SAMPLE and all("brak etykiety zdania" in i for i in issues)
    assert load_workbook(path)["Instrukcja"]["B2"].value == annotator
