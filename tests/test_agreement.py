"""Testy punktu 2: eksport anotacji i miary zgodności (kappa Cohena, kappa Fleissa)."""
import json

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from sklearn.metrics import cohen_kappa_score

import config
from src import agreement, export, sheets_io


# --- wzory referencyjne (liczone ręcznie, niezależnie od bibliotek) -------------

def manual_cohen(a, b) -> float:
    a, b = np.asarray(a), np.asarray(b)
    po = np.mean(a == b)
    pe = sum(np.mean(a == c) * np.mean(b == c) for c in np.union1d(a, b))
    return (po - pe) / (1 - pe)


def manual_fleiss(counts) -> float:
    """counts: (elementy, kategorie), w każdym wierszu ta sama liczba ocen."""
    counts = np.asarray(counts, dtype=float)
    n, N = counts.sum(axis=1)[0], len(counts)
    p_j = counts.sum(axis=0) / (N * n)
    P_i = (np.sum(counts ** 2, axis=1) - n) / (n * (n - 1))
    return (P_i.mean() - np.sum(p_j ** 2)) / (1 - np.sum(p_j ** 2))


def manual_scott(a, b) -> float:
    a, b = np.asarray(a), np.asarray(b)
    po = np.mean(a == b)
    pe = sum(((np.mean(a == c) + np.mean(b == c)) / 2) ** 2 for c in np.union1d(a, b))
    return (po - pe) / (1 - pe)


def _noisy_ratings(n=300, k=3, seed=0):
    """k oceniających: wspólna „prawdziwa” etykieta, każdy zmienia ją z prawdopodobieństwem 0,3."""
    rng = np.random.default_rng(seed)
    truth = rng.choice(config.SENT_LABELS, n, p=[0.3, 0.35, 0.25, 0.1])
    cols = [np.where(rng.random(n) < 0.3, rng.choice(config.SENT_LABELS, n), truth) for _ in range(k)]
    return np.column_stack(cols)


def _counts(ratings, categories):
    return np.array([[np.sum(row == c) for c in categories] for row in ratings])


# --- miary -----------------------------------------------------------------------

def test_cohen_matches_sklearn_and_manual_formula():
    r = _noisy_ratings(k=2)
    k = agreement.cohen(r[:, 0], r[:, 1])
    assert k == pytest.approx(cohen_kappa_score(r[:, 0], r[:, 1]))
    assert k == pytest.approx(manual_cohen(r[:, 0], r[:, 1]))
    assert 0.3 < k < 0.9


def test_fleiss_matches_manual_formula():
    r = _noisy_ratings(k=3)
    assert agreement.fleiss(r, config.SENT_LABELS) == pytest.approx(manual_fleiss(_counts(r, config.SENT_LABELS)))


def test_fleiss_wikipedia_example():
    """Przykład z artykułu „Fleiss' kappa” (Wikipedia): 10 elementów, 14 oceniających, 5 kategorii → 0,210."""
    counts = np.array([[0, 0, 0, 0, 14], [0, 2, 6, 4, 2], [0, 0, 3, 5, 6], [0, 3, 9, 2, 0], [2, 2, 8, 1, 1],
                       [7, 7, 0, 0, 0], [3, 2, 6, 3, 0], [2, 5, 3, 2, 2], [6, 5, 2, 1, 0], [0, 2, 2, 3, 7]])
    cats = ["1", "2", "3", "4", "5"]
    ratings = np.array([np.repeat(cats, row) for row in counts])
    assert agreement.fleiss(ratings, cats) == pytest.approx(0.210, abs=5e-4)
    assert manual_fleiss(counts) == pytest.approx(0.210, abs=5e-4)


def test_fleiss_with_two_raters_equals_scott_pi():
    r = _noisy_ratings(k=2, seed=1)
    assert agreement.fleiss(r, config.SENT_LABELS) == pytest.approx(manual_scott(r[:, 0], r[:, 1]))


def test_perfect_agreement_and_label_renaming():
    r = _noisy_ratings(k=2, seed=2)
    same = np.column_stack([r[:, 0], r[:, 0]])
    assert agreement.cohen(same[:, 0], same[:, 1]) == pytest.approx(1.0)
    assert agreement.fleiss(same, config.SENT_LABELS) == pytest.approx(1.0)
    assert agreement.raw_agreement(same) == 1.0
    rename = dict(zip(config.SENT_LABELS, ["w", "x", "y", "z"]))
    renamed = np.vectorize(rename.get)(r)
    assert agreement.cohen(renamed[:, 0], renamed[:, 1]) == pytest.approx(agreement.cohen(r[:, 0], r[:, 1]))
    assert agreement.fleiss(renamed, ["w", "x", "y", "z"]) == pytest.approx(agreement.fleiss(r, config.SENT_LABELS))


def test_fleiss_rejects_unknown_labels():
    with pytest.raises(ValueError):
        agreement.fleiss(np.array([["POZ", "XYZ"]]), config.SENT_LABELS)


@pytest.mark.parametrize("k, expected", [(-0.1, "poniżej losowej"), (0.0, "nieznaczna"), (0.2, "nieznaczna"),
                                         (0.35, "dostateczna"), (0.5, "umiarkowana"), (0.754, "znaczna"),
                                         (0.81, "prawie pełna"), (float("nan"), "–")])
def test_interpret_landis_koch(k, expected):
    assert agreement.interpret(k) == expected


def test_bootstrap_is_deterministic_and_brackets_point_estimate():
    r = _noisy_ratings(n=80, k=2, seed=3)
    lv = agreement.Level("sentences", ["A1", "A2"], r, np.arange(80), 80, config.SENT_LABELS)
    ci = agreement.bootstrap(lv, 300, seed=5)
    assert ci == agreement.bootstrap(lv, 300, seed=5)
    point = agreement.measures(lv, r)
    for m, (lo, hi) in ci.items():
        assert lo <= point[m] <= hi


def test_bootstrap_resamples_whole_sentences_for_tokens():
    # 2 zdania: pierwsze w pełni zgodne (3 tokeny), drugie w pełni niezgodne (1 token).
    r = np.array([["O", "O"], ["WZM", "WZM"], ["ODW", "ODW"], ["WZM", "ODW"]])
    lv = agreement.Level("tokens", ["A1", "A2"], r, np.array([0, 0, 0, 1]), 2, agreement.TOKEN_CATEGORIES)
    lo, hi = agreement.bootstrap(lv, 500, seed=0)["raw_agreement"]
    # losowanie całych zdań daje tylko wartości 0, 3/4 albo 1 – nigdy np. 1/2 czy 2/3
    assert lo == 0.0 and hi == 1.0


# --- eksport -----------------------------------------------------------------------

def test_token_spans_merge_and_offsets():
    text = "Lekarz bardzo miły , ale to strata czasu i nie polecam ."
    labels = ["O", "WZM", "WYR_POZ", "O", "O", "O", "WYR_NEG", "WYR_NEG", "O", "ODW", "WYR_POZ", "O"]
    spans = export.token_spans(text, labels)
    assert [s[2:] for s in spans] == [["WZM", "bardzo"], ["WYR_POZ", "miły"], ["WYR_NEG", "strata czasu"],
                                      ["ODW", "nie"], ["WYR_POZ", "polecam"]]
    assert all(text[s:e] == frag for s, e, _, frag in spans)
    with pytest.raises(ValueError):
        export.token_spans(text, labels[:-1])


TOY = pd.DataFrame({"sent_id": [11, 22], "text": ["Lekarz bardzo miły .", "Nie polecam tego hotelu ."]})


@pytest.fixture
def toy_iteration(tmp_path, monkeypatch):
    """Iteracja 1 w tmp: próbka 2 zdań, klucz gold, arkusze A1 i A2 (A2 bez etykiety jednego zdania)."""
    monkeypatch.setattr(config, "ANNOT_DIR", tmp_path / "annotation")
    d = config.iter_dir(1)
    d.mkdir(parents=True)
    TOY.to_csv(d / "sample.csv", index=False)
    pd.DataFrame({"sent_id": [11, 22], "gold": ["POZ", "NEG"], "gold_raw": ["z_plus_m", "z_minus_m"]}).to_csv(
        d / "_key_gold.csv", index=False)
    filled = config.filled_sheets_dir(1)
    filled.mkdir(parents=True)
    first_tok = len(sheets_io.FIXED_COLS) + 1
    fills = {"A1": {11: ("POZ", {1: "WZM", 2: "WYR_POZ"}), 22: ("NEG", {0: "ODW", 1: "WYR_POZ"})},
             "A2": {11: ("POZ", {2: "WYR_POZ"}), 22: (None, {0: "ODW", 1: "WYR_POZ"})}}
    for a, sentences in fills.items():
        wb = sheets_io.build_workbook(TOY, a, 1, "v1")
        ws = wb[sheets_io.SHEET]
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 2).value in sentences:
                label, toks = sentences[ws.cell(r, 2).value]
                ws.cell(r, 3, label)
                for j, lab in toks.items():
                    ws.cell(r + 1, first_tok + j, lab)
        wb.save(filled / f"iter1_{a}.xlsx")
    return filled


def test_export_stops_on_sheet_problems(toy_iteration):
    with pytest.raises(ValueError, match="iter1_A2.xlsx: .*brak etykiety zdania 22"):
        export.load_iteration(1)


def test_export_tables_on_toy_iteration(toy_iteration):
    wb_path = toy_iteration / "iter1_A2.xlsx"
    wb = load_workbook(wb_path)
    ws = wb[sheets_io.SHEET]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value == 22)
    ws.cell(row, 3, "NEU")
    wb.save(wb_path)

    paths = export.write_export(1)
    sent = pd.read_csv(paths["sentences_wide"])
    assert sent[["A1", "A2", "gold", "zgodne"]].values.tolist() == [["POZ", "POZ", "POZ", True], ["NEG", "NEU", "NEG", False]]
    tok = pd.read_csv(paths["tokens_wide"])
    assert len(tok) == 9 and tok["tok_idx"].tolist()[:4] == [1, 2, 3, 4]
    assert tok.loc[tok["sent_id"] == 11, "zgodne"].tolist() == [True, False, True, True]
    assert len(pd.read_csv(paths["tokens_long"])) == 18
    assert len(pd.read_csv(paths["sentences_long"])) == 4


# --- dane rzeczywiste (iteracja 1) --------------------------------------------------

REAL = config.filled_sheets_dir(1)


@pytest.fixture
def real_export(tmp_path, monkeypatch):
    if not (REAL / "iter1_A1.xlsx").exists():
        pytest.skip("brak wypełnionych arkuszy iteracji 1")
    monkeypatch.setattr(config, "export_dir", lambda it: tmp_path / f"export{it}")
    return export.write_export(1)


def test_real_export_matches_team_artefacts(real_export):
    """Eksport z kodu = ręczne zestawienia zespołu w iter1_artefacts/."""
    sent = pd.read_csv(real_export["sentences_wide"])
    team = pd.read_csv(REAL / "disagreements_sentences.csv")
    assert set(sent.loc[~sent["zgodne"], "sent_id"]) == set(team["sent_id"].astype(int))

    tok = pd.read_csv(real_export["tokens_wide"], keep_default_na=False)
    team_tok = pd.read_csv(REAL / "comparison_tokens.csv", keep_default_na=False)
    merged = tok.merge(team_tok, left_on=["sent_id", "tok_idx"], right_on=["sent_id", "pozycja_tokenu"],
                       suffixes=("", "_team"))
    assert len(merged) == len(tok) == len(team_tok) == 1706
    assert (merged["A1"] == merged["A1_team"]).all() and (merged["A2"] == merged["A2_team"]).all()
    assert tok["zgodne"].sum() == 1611


def test_real_spans_point_at_text(real_export):
    for a in ["A1", "A2"]:
        records = [json.loads(line) for line in open(real_export[f"spans_{a}"], encoding="utf-8")]
        assert len(records) == config.N_SAMPLE
        assert all(r["text"][s:e] == frag for r in records for s, e, _, frag in r["spans"])


def test_real_agreement_values(real_export):
    """Wartości z planu (PLAN_Punkt2_Zgodnosc.md, p. 6)."""
    res = agreement.compute(1, n_boot=50, seed=0)
    v = {lv: {m: x["value"] for m, x in d["measures"].items()} for lv, d in res["levels"].items()}
    assert v["sentences"]["raw_agreement"] == pytest.approx(0.83)
    assert v["sentences"]["cohen_A1-A2"] == pytest.approx(0.754, abs=5e-4)
    assert v["sentences"]["fleiss"] == pytest.approx(0.754, abs=5e-4)
    assert v["tokens"]["cohen_A1-A2"] == pytest.approx(0.730, abs=5e-4)
    assert v["tokens"]["fleiss"] == pytest.approx(0.730, abs=5e-4)
    assert res["levels"]["tokens_marked"]["n"] == 231
    assert v["tokens_marked"]["cohen_A1-A2"] == pytest.approx(0.463, abs=5e-4)
    assert v["tokens_marked"]["fleiss"] == pytest.approx(0.462, abs=5e-4)
    assert v["sentences_with_gold"]["fleiss"] == pytest.approx(0.636, abs=5e-4)
    assert sum(map(sum, res["confusion"]["sentences"])) == 100
    assert "π Scotta" in agreement.markdown_table(res)
